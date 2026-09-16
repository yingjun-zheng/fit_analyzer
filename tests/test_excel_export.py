"""Excel 报表导出 + 数据库备份 单测（P2-A）。

直接跑：python tests/test_excel_export.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core.db import DB

_results = {"ok": 0, "fail": 0}


def check(name, cond, extra=""):
    tag = "  [PASS]" if cond else "  [FAIL]"
    print(f"{tag} {name} {extra}")
    _results["ok" if cond else "fail"] += 1


def make_db(path, acts):
    """造库：acts = [(file_hash, name, start_time, dist_m, commute)]。"""
    db = DB(path)
    for i, (fh, name, st, dist, comm) in enumerate(acts, 1):
        db.conn.execute(
            """INSERT INTO activities (file_hash, name, start_time, start_ts, month,
               total_distance_m, timer_s, elapsed_s, moving_s, avg_speed_ms, max_speed_ms,
               avg_hr, avg_cad, calories, ascent_m, commute)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (fh, name, st, 1787904000 + i * 86400, st[:7], dist, 3600, 3700, 3500,
             4.2, 6.0, 150, 88, 500, 120, 1 if comm else 0))
    db.conn.commit()
    return db


def main():
    d = Path(tempfile.mkdtemp())
    db = make_db(d / "fit.db", [
        ("h1", "通勤-去程", "2026-09-01 08:00:00", 5000, True),
        ("h2", "周末长骑", "2026-09-06 07:30:00", 80000, False),
        ("h3", "通勤-返程", "2026-09-01 18:00:00", 5200, True),
        ("h4", "8月训练", "2026-08-15 07:00:00", 60000, False),
    ])

    from core import excel_export

    print("== Excel 报表导出（全部）==")
    out = d / "report.xlsx"
    n = excel_export.export_excel(db, out)
    check("返回活动条数 4", n == 4, str(n))
    check("文件已生成", out.exists())

    from openpyxl import load_workbook
    wb = load_workbook(out)
    check("Sheet 结构", wb.sheetnames == ["活动明细", "月度汇总"], str(wb.sheetnames))

    ws = wb["活动明细"]
    check("明细行数 = 4+表头", ws.max_row == 5, str(ws.max_row))
    rows = {r[0].value: r for r in ws.iter_rows(min_row=2)}
    tongqin = rows["2026-09-01 08:00"]
    check("通勤标记正确", tongqin[3].value == "通勤")
    check("距离口径 km", tongqin[4].value == 5.0, str(tongqin[4].value))
    check("8 月训练为训练类型", rows["2026-08-15 07:00"][3].value == "训练")

    ws2 = wb["月度汇总"]
    summary = {r[0].value: r for r in ws2.iter_rows(min_row=2)}
    sep = summary["2026-09"]
    check("9 月里程 = 90.2km", sep[2].value == 90.2, str(sep[2].value))
    check("9 月通勤 10.2km", sep[4].value == 10.2, str(sep[4].value))
    check("9 月训练 80.0km", sep[3].value == 80.0, str(sep[3].value))
    aug = summary["2026-08"]
    check("8 月通勤 0km", aug[4].value == 0, str(aug[4].value))

    print("== 按月导出 ==")
    out2 = d / "report_sep.xlsx"
    n2 = excel_export.export_excel(db, out2, month="2026-09")
    check("按月导出 3 条", n2 == 3, str(n2))
    wb2 = load_workbook(out2)
    check("按月明细行数", wb2["活动明细"].max_row == 4, str(wb2["活动明细"].max_row))

    print("== 空数据报错 ==")
    try:
        excel_export.export_excel(db, d / "empty.xlsx", month="2099-01")
        check("空月应抛错", False)
    except ValueError:
        check("空月应抛错", True)

    print("== 数据库备份 ==")
    dest = excel_export.backup_database(d / "fit.db", d / "backup")
    check("备份文件存在", dest.exists() and dest.name.startswith("fit_backup_"), dest.name)
    db2 = DB(dest)
    check("备份可打开且活动数一致", db2.count() == db.count())
    db2.close()

    db.close()
    total = _results["ok"] + _results["fail"]
    print(f"结果: {_results['ok']} 通过, {_results['fail']} 失败")
    sys.exit(1 if _results["fail"] else 0)


if __name__ == "__main__":
    main()
